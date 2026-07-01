#!/usr/bin/env python3
"""Build aggregate-only JSON for the public GLA GitHub Pages dashboard.

Input may contain private voter-level fields. Output must not.
"""
from __future__ import annotations
import csv, json, re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PRIVATE_SOURCE_DIR = ROOT / "private-source-data"
DEFAULT_INPUT = PRIVATE_SOURCE_DIR / "GLA_2026_Registration_Outreach_Tracking.csv"
LEGACY_INPUT = ROOT / "GLA_2026_Registration_Outreach_Tracking.csv"
ZODIAC_DIR = ROOT / "Zodiac Project 6.17.2026"
ZODIAC_STATEWIDE_INPUT = PRIVATE_SOURCE_DIR / "Registered_voters_by_Zodiac.xlsx"
ZODIAC_PULASKI_INPUT = PRIVATE_SOURCE_DIR / "Registered_voters_in_Pulaski_by_Zodiac.xlsx"
LEGACY_ZODIAC_STATEWIDE_INPUT = ZODIAC_DIR / "Registered_voters_by_Zodiac.xlsx"
LEGACY_ZODIAC_PULASKI_INPUT = ZODIAC_DIR / "Registered_voters_in_Pulaski_by_Zodiac.xlsx"
DEFAULT_OUTPUT = ROOT / "github-pages" / "data" / "public-dashboard.json"
PRIVATE_TERMS = re.compile(r"first name|last name|voterid|voter id|full address|phone number|email address|birth date|birthday", re.I)
ZODIAC_ORDER = [
    "Aries", "Taurus", "Gemini", "Cancer", "Leo", "Virgo",
    "Libra", "Scorpio", "Sagittarius", "Capricorn", "Aquarius", "Pisces",
]

def clean(value: str | None) -> str:
    value = (value or "").strip()
    return value if value else "Unknown"

def default_input_path() -> Path:
    if DEFAULT_INPUT.exists():
        return DEFAULT_INPUT
    if LEGACY_INPUT.exists():
        return LEGACY_INPUT
    return DEFAULT_INPUT

def first_existing(*paths: Path) -> Path:
    for path in paths:
        if path.exists():
            return path
    return paths[0]

def safe_int(value: Any) -> int:
    if value is None or value == "":
        return 0
    try:
        return int(float(str(value).replace(",", "")))
    except (TypeError, ValueError):
        return 0

def pct(numerator: float, denominator: float) -> float:
    return round((numerator / denominator * 100), 2) if denominator else 0

def modified_at(path: Path) -> str | None:
    if not path.exists():
        return None
    return datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()

def sum_key(rows: list[dict[str, Any]], key: str) -> float:
    return sum(float(row.get(key) or 0) for row in rows)

def read_zodiac_rows(path: Path, county_column: str | None) -> list[dict[str, Any]]:
    if not path.exists():
        return []

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    rows = []

    for raw_row in sheet.iter_rows(min_row=2, values_only=True):
        raw = dict(zip(headers, raw_row))
        sign = raw.get("z.Zodiac Sign")
        if not sign:
            continue
        row = {
            "zodiacSign": str(sign),
            "activeRegisteredVoters": safe_int(raw.get("Active Registered Voters")),
            "activeRegisteredPct": float(raw.get("% of Active Registered Voters") or raw.get("% of County Active Registered Voters") or 0),
            "recentlyVoted": safe_int(raw.get("Recently Voted")),
            "recentlyVotedPct": float(raw.get("% of Zodiac That Recently Voted") or 0),
            "didNotRecentlyVote": safe_int(raw.get("Did Not Recently Vote")),
        }
        if county_column:
            row["county"] = clean(raw.get(county_column))
        rows.append(row)
    return rows

def aggregate_zodiac_by_sign(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = {}
    total_active = sum_key(rows, "activeRegisteredVoters")
    for row in rows:
        sign = str(row["zodiacSign"])
        item = grouped.setdefault(
            sign,
            {
                "zodiacSign": sign,
                "activeRegisteredVoters": 0,
                "recentlyVoted": 0,
                "didNotRecentlyVote": 0,
            },
        )
        item["activeRegisteredVoters"] += int(row["activeRegisteredVoters"])
        item["recentlyVoted"] += int(row["recentlyVoted"])
        item["didNotRecentlyVote"] += int(row["didNotRecentlyVote"])

    output = []
    for item in grouped.values():
        active = float(item["activeRegisteredVoters"])
        recent = float(item["recentlyVoted"])
        item["activeRegisteredPct"] = pct(active, total_active)
        item["recentlyVotedPct"] = pct(recent, active)
        output.append(item)
    return sorted(output, key=lambda row: ZODIAC_ORDER.index(str(row["zodiacSign"])) if str(row["zodiacSign"]) in ZODIAC_ORDER else 99)

def load_zodiac_data() -> dict[str, Any] | None:
    statewide_path = first_existing(ZODIAC_STATEWIDE_INPUT, LEGACY_ZODIAC_STATEWIDE_INPUT)
    pulaski_path = first_existing(ZODIAC_PULASKI_INPUT, LEGACY_ZODIAC_PULASKI_INPUT)
    statewide_rows = read_zodiac_rows(statewide_path, "z.County")
    pulaski_rows = read_zodiac_rows(pulaski_path, None)
    if not statewide_rows and not pulaski_rows:
        return None

    statewide_by_sign = aggregate_zodiac_by_sign(statewide_rows)
    county_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in statewide_rows:
        county_groups[str(row.get("county") or "Unknown")].append(row)

    county_summary = []
    for county, rows in county_groups.items():
        active = sum_key(rows, "activeRegisteredVoters")
        recent = sum_key(rows, "recentlyVoted")
        county_summary.append(
            {
                "county": county,
                "activeRegisteredVoters": int(active),
                "recentlyVoted": int(recent),
                "didNotRecentlyVote": int(sum_key(rows, "didNotRecentlyVote")),
                "recentlyVotedPct": pct(recent, active),
                "topZodiacByRecentlyVotedPct": max(rows, key=lambda row: float(row["recentlyVotedPct"]))["zodiacSign"] if rows else None,
            }
        )

    county_summary.sort(key=lambda row: (-float(row["recentlyVotedPct"]), str(row["county"])))
    pulaski_rows.sort(key=lambda row: -float(row["recentlyVotedPct"]))
    active_total = sum_key(statewide_rows, "activeRegisteredVoters")
    recent_total = sum_key(statewide_rows, "recentlyVoted")

    return {
        "source": {
            "statewideFileName": statewide_path.name if statewide_path.exists() else None,
            "statewideModifiedAt": modified_at(statewide_path),
            "pulaskiFileName": pulaski_path.name if pulaski_path.exists() else None,
            "pulaskiModifiedAt": modified_at(pulaski_path),
        },
        "summary": {
            "activeRegisteredVoters": int(active_total),
            "recentlyVoted": int(recent_total),
            "didNotRecentlyVote": int(sum_key(statewide_rows, "didNotRecentlyVote")),
            "recentlyVotedPct": pct(recent_total, active_total),
            "countyCount": len(county_summary),
            "zodiacSigns": len(statewide_by_sign),
            "pulaskiRows": len(pulaski_rows),
        },
        "statewideBySign": statewide_by_sign,
        "countySummary": county_summary,
        "pulaskiBySign": pulaski_rows,
    }

def main(input_path: str | None = None, output_path: str | None = None) -> None:
    in_path = Path(input_path) if input_path else default_input_path()
    out_path = Path(output_path) if output_path else DEFAULT_OUTPUT
    if not in_path.exists():
        raise SystemExit(
            "Cannot find source CSV. Put the latest file at "
            f"{DEFAULT_INPUT} or pass --input explicitly."
        )
    rows = []
    with in_path.open(newline="", encoding="utf-8-sig", errors="replace") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(row)

    total = len(rows)
    counties = defaultdict(lambda: Counter(people=0, addedToCurrentVrvh=0, contacted=0, active=0, inactive=0))
    sources = defaultdict(lambda: Counter(people=0, addedToCurrentVrvh=0, notInGlaContactFile=0))
    reg_status = Counter()
    change_status = Counter()
    removed_status = Counter()
    contact_status = Counter()

    for row in rows:
        county = clean(row.get("c.County"))
        status = clean(row.get("Registration Status"))
        change = clean(row.get("Change Status"))
        removed = clean(row.get("Removed Tracking Status"))
        contact = clean(row.get("GLA Contact Status"))
        source = clean(row.get("Updated GLA Contact Source"))
        added = change.upper() == "ADDED TO CURRENT VRVH"
        contacted = contact.upper() != "NOT IN GLA CONTACT FILE"
        active = status.upper() == "A"
        inactive = status.upper() == "I"

        counties[county]["people"] += 1
        counties[county]["addedToCurrentVrvh"] += int(added)
        counties[county]["contacted"] += int(contacted)
        counties[county]["active"] += int(active)
        counties[county]["inactive"] += int(inactive)
        sources[source]["people"] += 1
        sources[source]["addedToCurrentVrvh"] += int(added)
        sources[source]["notInGlaContactFile"] += int(contact.upper() == "NOT IN GLA CONTACT FILE")
        reg_status[status] += 1
        change_status[change] += 1
        removed_status[removed] += 1
        contact_status[contact] += 1

    added_total = sum(1 for r in rows if clean(r.get("Change Status")).upper() == "ADDED TO CURRENT VRVH")
    contacted_total = sum(1 for r in rows if clean(r.get("GLA Contact Status")).upper() != "NOT IN GLA CONTACT FILE")
    active_total = sum(1 for r in rows if clean(r.get("Registration Status")).upper() == "A")
    inactive_total = sum(1 for r in rows if clean(r.get("Registration Status")).upper() == "I")
    not_in_gla = contact_status.get("NOT IN GLA CONTACT FILE", 0)

    def counter_rows(counter: Counter, label: str):
        return [{label: k, "count": v} for k, v in counter.most_common()]
    def county_rows():
        return [{"county": k, **dict(v)} for k, v in sorted(counties.items(), key=lambda item: (-item[1]["people"], item[0]))]
    def source_rows():
        return [{"source": k, **dict(v)} for k, v in sorted(sources.items(), key=lambda item: (-item[1]["people"], item[0]))]

    data = {
        "meta": {"title": "Get Loud Arkansas Public Dashboard", "generatedAt": datetime.now(timezone.utc).date().isoformat(), "privacy": "aggregate-only"},
        "overview": {"totalPeople": total, "addedToCurrentVrvh": added_total, "contacted": contacted_total, "notInGlaContactFile": not_in_gla, "counties": len(counties), "contactRate": round((contacted_total / total * 100), 1) if total else 0},
        "tracking": {"summary": {"totalPeople": total, "addedToCurrentVrvh": added_total, "notInGlaContactFile": not_in_gla, "counties": len(counties)}, "sourceTotals": source_rows(), "countyTotals": county_rows()},
        "voterFile": {"summary": {"totalPeople": total, "active": active_total, "inactive": inactive_total}, "registrationStatus": counter_rows(reg_status, "status")},
        "purge": {"summary": {"totalPeople": total, "addedToCurrentVrvh": added_total}, "removedTrackingStatus": counter_rows(removed_status, "status")},
        "outreach": {"summary": {"totalPeople": total, "contacted": contacted_total, "notInGlaContactFile": not_in_gla, "contactRate": round((contacted_total / total * 100), 1) if total else 0}, "contactStatus": counter_rows(contact_status, "status")},
        "registration": {"summary": {"totalPeople": total, "addedToCurrentVrvh": added_total, "active": active_total, "inactive": inactive_total}, "changeStatus": counter_rows(change_status, "status")},
        "geography": {"counties": county_rows()}
    }
    zodiac = load_zodiac_data()
    if zodiac:
        data["zodiac"] = zodiac

    encoded = json.dumps(data, indent=2, ensure_ascii=False)
    if PRIVATE_TERMS.search(encoded):
        raise SystemExit("Refusing to write public JSON because private field names were detected.")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(encoded + "\n", encoding="utf-8")
    print(f"Wrote {out_path} from {total:,} private source rows; output is aggregate-only.")

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default=None)
    parser.add_argument("--output", default=None)
    args = parser.parse_args()
    main(args.input, args.output)
